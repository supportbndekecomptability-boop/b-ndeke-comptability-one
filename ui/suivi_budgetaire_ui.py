"""
Page Suivi budgetaire - Prevu / Reel / Ecart
Avec export Excel + impression PDF.
"""
import os
from datetime import datetime
import customtkinter as ctk
from tkinter import messagebox, filedialog
from config import COLOR_NAVY, format_montant, DATA_DIR
from core.depenses import suivi_budgetaire_complet
from core.budgets import liste_budgets


class SuiviBudgetaireWindow(ctk.CTkToplevel):
    def __init__(self, parent):
        super().__init__(parent)

        self.title("Suivi budgetaire")
        largeur = min(1050, self.winfo_screenwidth() - 60)
        hauteur = min(760, self.winfo_screenheight() - 100)

        self.geometry(f"{largeur}x{hauteur}")
        self.configure(fg_color="#F5F7FB")
        self.resizable(False, False)

        x = (self.winfo_screenwidth() // 2) - (largeur // 2)
        y = max(10, (self.winfo_screenheight() // 2) - (hauteur // 2) - 20)
        self.geometry(f"{largeur}x{hauteur}+{x}+{y}")

        self.bind("<Escape>", lambda e: self.destroy())
        self.transient(parent)
        self.after(100, self._activer_grab)

        # Charger les donnees
        try:
            self.suivi = suivi_budgetaire_complet()
        except Exception as e:
            self.suivi = {}
            print(f"[SUIVI] Erreur : {e}")

        self._construire()

    def _activer_grab(self):
        try:
            self.grab_set()
            self.focus_force()
        except Exception:
            pass

    def _construire(self):
        # ===== HEADER =====
        header = ctk.CTkFrame(self, fg_color=COLOR_NAVY, corner_radius=0, height=70)
        header.pack(fill="x")
        header.pack_propagate(False)

        ctk.CTkLabel(
            header, text="SUIVI BUDGETAIRE",
            font=("Segoe UI", 16, "bold"),
            text_color="white",
        ).pack(side="left", padx=20, pady=20)

        ctk.CTkButton(
            header, text="Fermer",
            font=("Segoe UI", 11),
            fg_color="#e74c3c", hover_color="#c0392b",
            width=90, height=34,
            command=self.destroy,
        ).pack(side="right", padx=(5, 20), pady=18)

        ctk.CTkButton(
            header, text="Imprimer (PDF)",
            font=("Segoe UI", 11, "bold"),
            fg_color="#3498db", hover_color="#2980b9",
            width=130, height=34,
            command=self._imprimer_pdf,
        ).pack(side="right", padx=5, pady=18)

        ctk.CTkButton(
            header, text="Exporter Excel",
            font=("Segoe UI", 11, "bold"),
            fg_color="#27ae60", hover_color="#229954",
            width=130, height=34,
            command=self._exporter_excel,
        ).pack(side="right", padx=5, pady=18)

        # ===== CORPS =====
        scroll = ctk.CTkScrollableFrame(self, fg_color="transparent")
        scroll.pack(fill="both", expand=True, padx=15, pady=15)

        if not self.suivi:
            ctk.CTkLabel(
                scroll,
                text="Aucun budget configure.\n\n"
                     "Allez dans Parametres > Budgets pour definir vos postes.",
                font=("Segoe UI", 12),
                text_color="#666666",
                justify="center",
            ).pack(pady=50)
            return

        self._afficher_global(scroll)

        for key, label, couleur in liste_budgets():
            b = self.suivi.get(key, {})
            if b:
                self._afficher_budget(scroll, label, couleur, b)

    def _afficher_global(self, parent):
        cadre = ctk.CTkFrame(parent, fg_color="#0F2C5C", corner_radius=10)
        cadre.pack(fill="x", pady=(0, 15))

        ctk.CTkLabel(
            cadre, text="RECAPITULATIF GLOBAL",
            font=("Segoe UI", 12, "bold"),
            text_color="#F2B705",
        ).pack(pady=(12, 8))

        ligne = ctk.CTkFrame(cadre, fg_color="transparent")
        ligne.pack(fill="x", padx=15, pady=(0, 12))

        total_prevu = sum(self.suivi[k]["prevu"] for k in self.suivi)
        total_reel = sum(self.suivi[k]["reel"] for k in self.suivi)
        total_ecart = total_prevu - total_reel

        couleur_ecart = "#27ae60" if total_ecart >= 0 else "#e74c3c"
        symbole = "+" if total_ecart >= 0 else ""

        for titre, valeur, couleur in [
            ("Total prevu", format_montant(total_prevu), "#F2B705"),
            ("Total depense", format_montant(total_reel), "#e74c3c"),
            ("Ecart", f"{symbole}{format_montant(total_ecart)}", couleur_ecart),
        ]:
            bloc = ctk.CTkFrame(ligne, fg_color="transparent")
            bloc.pack(side="left", expand=True)

            ctk.CTkLabel(
                bloc, text=titre,
                font=("Segoe UI", 10),
                text_color="#cccccc",
            ).pack()

            ctk.CTkLabel(
                bloc, text=valeur,
                font=("Segoe UI", 16, "bold"),
                text_color=couleur,
            ).pack(pady=(3, 0))

    def _afficher_budget(self, parent, label, couleur, data):
        carte = ctk.CTkFrame(parent, fg_color="white", corner_radius=10,
                             border_width=2, border_color=couleur)
        carte.pack(fill="x", pady=(0, 12))

        entete = ctk.CTkFrame(carte, fg_color=couleur, corner_radius=8)
        entete.pack(fill="x", padx=8, pady=(8, 8))

        ctk.CTkLabel(
            entete, text=label,
            font=("Segoe UI", 13, "bold"),
            text_color="white",
        ).pack(side="left", padx=12, pady=8)

        prevu = data["prevu"]
        reel = data["reel"]
        ecart = data["ecart"]

        texte_ecart = self._format_ecart(ecart)
        couleur_ecart = "#27ae60" if ecart >= 0 else "#e74c3c"

        ctk.CTkLabel(
            entete,
            text=f"Prevu {format_montant(prevu)} | Reel {format_montant(reel)} | Ecart ",
            font=("Segoe UI", 10),
            text_color="white",
        ).pack(side="right", padx=(0, 4), pady=8)

        ctk.CTkLabel(
            entete, text=texte_ecart,
            font=("Segoe UI", 11, "bold"),
            text_color=couleur_ecart,
        ).pack(side="right", padx=(0, 12), pady=8)

        # Barre
        if prevu > 0:
            pct = min(100, (reel / prevu) * 100)
        else:
            pct = 0

        bar_frame = ctk.CTkFrame(carte, fg_color="#f0f0f0", height=8, corner_radius=4)
        bar_frame.pack(fill="x", padx=15, pady=(5, 3))
        bar_frame.pack_propagate(False)

        couleur_barre = "#e74c3c" if pct > 100 else ("#e67e22" if pct > 80 else "#27ae60")
        ctk.CTkFrame(bar_frame, fg_color=couleur_barre, corner_radius=4,
                     height=8).place(x=0, y=0, relwidth=pct / 100.0, relheight=1)

        ctk.CTkLabel(
            carte, text=f"{pct:.0f}% utilise",
            font=("Segoe UI", 9),
            text_color="#888888",
        ).pack(anchor="e", padx=15, pady=(0, 8))

        # Entete sous-categories
        entete_sous = ctk.CTkFrame(carte, fg_color="#f0f0f0", corner_radius=6)
        entete_sous.pack(fill="x", padx=12, pady=(0, 4))

        for nom, larg in [
            ("Sous-categorie", 300),
            ("Prevu", 130),
            ("Reel", 130),
            ("Ecart", 180),
        ]:
            ctk.CTkLabel(
                entete_sous, text=nom,
                font=("Segoe UI", 11, "bold"),
                text_color=COLOR_NAVY,
                width=larg, anchor="w",
            ).pack(side="left", padx=4, pady=8)

        for nom, info in data["sous_categories"].items():
            self._afficher_ligne_sous(carte, nom, info)

        ctk.CTkLabel(carte, text="").pack(pady=4)

    def _afficher_ligne_sous(self, parent, nom, info):
        prevu = info["prevu"]
        reel = info["reel"]
        ecart = info["ecart"]

        couleur_ecart = "#27ae60" if ecart >= 0 else "#e74c3c"
        texte_ecart = self._format_ecart(ecart)

        ligne = ctk.CTkFrame(parent, fg_color="#fafafa", corner_radius=4)
        ligne.pack(fill="x", padx=12, pady=2)

        ctk.CTkLabel(
            ligne, text=nom,
            font=("Segoe UI", 11),
            text_color="#333333",
            width=300, anchor="w",
        ).pack(side="left", padx=4, pady=6)

        ctk.CTkLabel(
            ligne, text=format_montant(prevu),
            font=("Segoe UI", 11, "bold"),
            text_color="#8e44ad",
            width=130, anchor="w",
        ).pack(side="left", padx=4, pady=6)

        ctk.CTkLabel(
            ligne, text=format_montant(reel),
            font=("Segoe UI", 11, "bold"),
            text_color="#e74c3c",
            width=130, anchor="w",
        ).pack(side="left", padx=4, pady=6)

        indicateur = ""
        if reel > prevu:
            indicateur = "  DEPASSEMENT"
        elif reel > prevu * 0.9:
            indicateur = "  ATTENTION"

        ctk.CTkLabel(
            ligne, text=texte_ecart + indicateur,
            font=("Segoe UI", 11, "bold"),
            text_color=couleur_ecart,
            width=180, anchor="w",
        ).pack(side="left", padx=4, pady=6)

    def _format_ecart(self, ecart):
        if ecart >= 0:
            return f"+ {format_montant(ecart)}"
        else:
            return f"- {format_montant(abs(ecart))}"

    # =========================================================
    # EXPORT EXCEL
    # =========================================================
    def _exporter_excel(self):
        if not self.suivi:
            messagebox.showwarning("Vide", "Aucune donnee a exporter.")
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            messagebox.showerror(
                "Erreur",
                "La bibliotheque 'openpyxl' n'est pas installee.\n\n"
                "Lancez : pip install openpyxl"
            )
            return

        defaut = f"Suivi_budgetaire_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        chemin = filedialog.asksaveasfilename(
            title="Enregistrer le suivi budgetaire",
            defaultextension=".xlsx",
            initialfile=defaut,
            filetypes=[("Excel", "*.xlsx"), ("Tous les fichiers", "*.*")],
        )
        if not chemin:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Suivi budgetaire"

            titre_font = Font(bold=True, size=14, color="FFFFFF")
            titre_fill = PatternFill("solid", fgColor="0F2C5C")
            entete_font = Font(bold=True, color="FFFFFF")
            entete_fill = PatternFill("solid", fgColor="0F2C5C")
            total_font = Font(bold=True)
            total_fill = PatternFill("solid", fgColor="E8F4EA")
            sous_fill = PatternFill("solid", fgColor="F5F5F5")
            centre = Alignment(horizontal="center", vertical="center")
            droite = Alignment(horizontal="right", vertical="center")
            gauche = Alignment(horizontal="left", vertical="center")
            bord = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"),
            )

            ligne = 1

            ws.cell(row=ligne, column=1, value="SUIVI BUDGETAIRE").font = titre_font
            ws.cell(row=ligne, column=1).fill = titre_fill
            ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
            ws.cell(row=ligne, column=1).alignment = centre
            ws.row_dimensions[ligne].height = 25
            ligne += 1

            ws.cell(row=ligne, column=1,
                    value=f"Genere le {datetime.now().strftime('%d/%m/%Y %H:%M')}")
            ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
            ws.cell(row=ligne, column=1).alignment = centre
            ligne += 2

            # RECAP GLOBAL
            total_prevu = sum(self.suivi[k]["prevu"] for k in self.suivi)
            total_reel = sum(self.suivi[k]["reel"] for k in self.suivi)
            total_ecart = total_prevu - total_reel

            ws.cell(row=ligne, column=1, value="RECAPITULATIF GLOBAL").font = entete_font
            ws.cell(row=ligne, column=1).fill = entete_fill
            ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
            ws.cell(row=ligne, column=1).alignment = centre
            ligne += 1

            for col, val in enumerate(["Total prevu", "Total depense", "Ecart"], start=1):
                c = ws.cell(row=ligne, column=col, value=val)
                c.font = Font(bold=True)
                c.fill = sous_fill
                c.alignment = centre
                c.border = bord
            ligne += 1

            for col, val in enumerate(
                [total_prevu, total_reel,
                 ("+" if total_ecart >= 0 else "-") + str(abs(total_ecart))],
                start=1,
            ):
                c = ws.cell(row=ligne, column=col, value=val)
                c.font = total_font
                c.fill = total_fill
                c.alignment = droite if col <= 2 else centre
                c.border = bord
            ligne += 2

            # DETAIL
            labels = {
                "prime": "PRIME",
                "investissement": "INVESTISSEMENT",
                "fonctionnement": "FONCTIONNEMENT",
            }

            for key, label in labels.items():
                b = self.suivi.get(key, {})
                if not b:
                    continue

                ws.cell(row=ligne, column=1, value=label).font = entete_font
                ws.cell(row=ligne, column=1).fill = entete_fill
                ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=4)
                ws.cell(row=ligne, column=1).alignment = centre
                ligne += 1

                for col, val in enumerate(
                    ["Prevu", "Reel", "Ecart", "Utilise %"], start=1
                ):
                    c = ws.cell(row=ligne, column=col, value=val)
                    c.font = Font(bold=True)
                    c.fill = sous_fill
                    c.alignment = centre
                    c.border = bord
                ligne += 1

                p = b["prevu"]
                r = b["reel"]
                e = b["ecart"]
                pct = min(100, (r / p * 100) if p > 0 else 0)

                for col, val in enumerate([p, r, e, f"{pct:.0f}%"], start=1):
                    c = ws.cell(row=ligne, column=col, value=val)
                    c.alignment = centre if col == 4 else droite
                    c.border = bord
                ligne += 1

                for col, val in enumerate(
                    ["Sous-categorie", "Prevu", "Reel", "Ecart"], start=1
                ):
                    c = ws.cell(row=ligne, column=col, value=val)
                    c.font = Font(bold=True)
                    c.fill = sous_fill
                    c.alignment = centre if col != 1 else gauche
                    c.border = bord
                ligne += 1

                for nom, info in b["sous_categories"].items():
                    ws.cell(row=ligne, column=1, value=nom).border = bord
                    ws.cell(row=ligne, column=1).alignment = gauche

                    ws.cell(row=ligne, column=2, value=info["prevu"]).alignment = droite
                    ws.cell(row=ligne, column=2).border = bord

                    ws.cell(row=ligne, column=3, value=info["reel"]).alignment = droite
                    ws.cell(row=ligne, column=3).border = bord

                    ws.cell(row=ligne, column=4, value=info["ecart"]).alignment = droite
                    ws.cell(row=ligne, column=4).border = bord
                    ligne += 1

                ligne += 1

            ws.column_dimensions["A"].width = 42
            ws.column_dimensions["B"].width = 18
            ws.column_dimensions["C"].width = 18
            ws.column_dimensions["D"].width = 18

            wb.save(chemin)

            messagebox.showinfo(
                "Export Excel",
                f"Fichier enregistre :\n{chemin}"
            )

            try:
                os.startfile(chemin)
            except Exception:
                pass

        except Exception as e:
            import traceback
            traceback.print_exc()
            messagebox.showerror("Erreur", f"Erreur lors de l'export : {e}")

    # =========================================================
    # IMPRIMER (PDF)
    # =========================================================
    def _imprimer_pdf(self):
        if not self.suivi:
            messagebox.showwarning("Vide", "Aucune donnee a imprimer.")
            return

        try:
            from services.etats_pdf import pdf_suivi_budgetaire
            from ui.pdf_viewer import PdfViewerWindow

            chemin = pdf_suivi_budgetaire(self.suivi)
            PdfViewerWindow(self, chemin, titre="Suivi budgetaire")

        except Exception as e:
            import traceback
            traceback.print_exc()
            messagebox.showerror("Erreur", f"Erreur : {e}")