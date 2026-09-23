"""
Export des donnees vers Excel (.xlsx) avec plusieurs onglets
"""
import os
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False

from config import DATA_DIR, format_montant
from database import get_connection
from core.etats_financiers import (
    etat_journal, etat_balance, etat_bilan, etat_compte_resultat,
)

EXPORTS_DIR = os.path.join(DATA_DIR, "exports")
os.makedirs(EXPORTS_DIR, exist_ok=True)


# Couleurs (OpenPyXL veut du ARGB sans #)
NAVY_HEX = "FF0F2C5C"
GOLD_HEX = "FFF2B705"
WHITE_HEX = "FFFFFFFF"
GREY_HEX = "FFF0F0F0"


def _creer_onglet(wb, nom_onglet, titre, entetes, lignes, largeurs=None):
    """
    Cree un onglet dans le classeur Excel avec :
    - Un titre en haut (ligne 1)
    - Un en-tete (ligne 2)
    - Les donnees (lignes 3+)
    """
    ws = wb.create_sheet(title=nom_onglet[:31])  # Excel limite a 31 car

    # Ligne 1 : Titre principal
    if entetes:
        ws.merge_cells(start_row=1, start_column=1,
                       end_row=1, end_column=len(entetes))
    ws.cell(row=1, column=1, value=titre)
    ws.cell(row=1, column=1).font = Font(bold=True, size=14, color=NAVY_HEX)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    # Ligne 2 : En-tetes
    for col_idx, entete in enumerate(entetes, start=1):
        c = ws.cell(row=2, column=col_idx, value=entete)
        c.font = Font(bold=True, color=WHITE_HEX)
        c.fill = PatternFill("solid", fgColor=NAVY_HEX)
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Lignes de donnees
    for row_idx, ligne in enumerate(lignes, start=3):
        for col_idx, val in enumerate(ligne, start=1):
            c = ws.cell(row=row_idx, column=col_idx, value=val)
            # Alterner les couleurs
            if row_idx % 2 == 0:
                c.fill = PatternFill("solid", fgColor="FFFAFAFA")

    # Largeurs de colonnes
    if largeurs:
        for col_idx, largeur in enumerate(largeurs, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = largeur
    else:
        for col_idx in range(1, len(entetes) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 18

    return ws


def _formater_nombre(ws, colonne, debut_ligne=3, nb_decimales=0):
    """Formate une colonne en nombre avec separateur de milliers"""
    for row in range(debut_ligne, ws.max_row + 1):
        c = ws.cell(row=row, column=colonne)
        if c.value is not None and isinstance(c.value, (int, float)):
            c.number_format = '#,##0' if nb_decimales == 0 else '#,##0.00'


def exporter_tout_vers_excel():
    """
    Cree un fichier Excel avec TOUS les onglets.
    Retourne (ok, chemin_ou_erreur)
    """
    if not OPENPYXL_OK:
        return False, ("Module openpyxl non installe.\n\n"
                       "Executez : python -m pip install openpyxl")

    try:
        wb = openpyxl.Workbook()
        # Supprimer la feuille par defaut
        wb.remove(wb.active)

        conn = get_connection()
        cursor = conn.cursor()

        # ============ 1. ELEVES ============
        cursor.execute("""
            SELECT e.matricule, e.nom, e.prenom, e.classe, e.sexe,
                   e.nom_parent, e.telephone_parent,
                   COALESCE(e.frais_scolarite, 0) as frais,
                   COALESCE(e.frais_scolarite, 0) - COALESCE(
                       (SELECT SUM(montant) FROM paiements_eleves
                        WHERE eleve_id = e.id), 0
                   ) as solde
            FROM eleves e
            WHERE e.actif = 1
            ORDER BY e.classe, e.nom, e.prenom
        """)
        eleves = cursor.fetchall()

        lignes = []
        for e in eleves:
            lignes.append([
                e["matricule"], e["nom"], e["prenom"], e["classe"],
                e["sexe"] or "", e["nom_parent"] or "",
                e["telephone_parent"] or "",
                float(e["frais"] or 0), float(e["solde"] or 0),
            ])

        _creer_onglet(
            wb, "Eleves", "LISTE DES ELEVES",
            ["Matricule", "Nom", "Prenom", "Classe", "Sexe",
             "Parent", "Telephone", "Frais", "Solde"],
            lignes,
            largeurs=[15, 18, 18, 14, 10, 20, 16, 14, 14],
        )
        ws = wb["Eleves"]
        _formater_nombre(ws, 8)
        _formater_nombre(ws, 9)

        # ============ 2. PERSONNEL ============
        cursor.execute("""
            SELECT p.code, p.nom, p.prenom, p.fonction, p.sexe,
                   p.telephone, p.salaire_mensuel, p.duree_contrat_mois,
                   COALESCE(p.salaire_mensuel, 0) * COALESCE(p.duree_contrat_mois, 12)
                       as engagement,
                   COALESCE((SELECT SUM(montant) FROM paiements_personnel
                             WHERE personnel_id = p.id), 0) as paye
            FROM personnel p
            WHERE p.actif = 1
            ORDER BY p.nom, p.prenom
        """)
        personnel = cursor.fetchall()

        lignes = []
        for p in personnel:
            engagement = float(p["engagement"] or 0)
            paye = float(p["paye"] or 0)
            lignes.append([
                p["code"], p["nom"], p["prenom"], p["fonction"],
                p["sexe"] or "", p["telephone"] or "",
                float(p["salaire_mensuel"] or 0), p["duree_contrat_mois"],
                engagement, paye, engagement - paye,
            ])

        _creer_onglet(
            wb, "Personnel", "LISTE DU PERSONNEL",
            ["Code", "Nom", "Prenom", "Fonction", "Sexe", "Telephone",
             "Salaire/mois", "Duree (mois)", "Engagement", "Paye", "Reste"],
            lignes,
            largeurs=[14, 18, 18, 18, 10, 16, 14, 12, 14, 14, 14],
        )
        ws = wb["Personnel"]
        for col in (7, 9, 10, 11):
            _formater_nombre(ws, col)

        # ============ 3. PAIEMENTS ELEVES ============
        cursor.execute("""
            SELECT p.numero_recu, p.date_paiement, e.matricule,
                   e.nom, e.prenom, e.classe, p.montant, p.motif, p.mode_paiement
            FROM paiements_eleves p
            JOIN eleves e ON p.eleve_id = e.id
            ORDER BY p.date_paiement DESC
        """)
        paiements_eleves = cursor.fetchall()

        lignes = []
        for p in paiements_eleves:
            lignes.append([
                p["numero_recu"],
                p["date_paiement"][:19] if p["date_paiement"] else "",
                p["matricule"], p["nom"], p["prenom"], p["classe"],
                float(p["montant"] or 0), p["motif"], p["mode_paiement"],
            ])

        _creer_onglet(
            wb, "Paiements Eleves", "PAIEMENTS DES ELEVES",
            ["N Recu", "Date", "Matricule", "Nom", "Prenom", "Classe",
             "Montant", "Motif", "Mode"],
            lignes,
            largeurs=[16, 20, 14, 16, 16, 12, 14, 20, 16],
        )
        _formater_nombre(wb["Paiements Eleves"], 7)

        # ============ 4. PAIEMENTS PERSONNEL ============
        cursor.execute("""
            SELECT p.numero_paie, p.date_paiement, e.code, e.nom, e.prenom,
                   e.fonction, p.montant,
                   COALESCE(p.montant_avance_deduit, 0) as avance,
                   p.motif, p.mode_paiement
            FROM paiements_personnel p
            JOIN personnel e ON p.personnel_id = e.id
            ORDER BY p.date_paiement DESC
        """)
        paiements_pers = cursor.fetchall()

        lignes = []
        for p in paiements_pers:
            m = float(p["montant"] or 0)
            a = float(p["avance"] or 0)
            lignes.append([
                p["numero_paie"],
                p["date_paiement"][:19] if p["date_paiement"] else "",
                p["code"], p["nom"], p["prenom"], p["fonction"],
                m, a, m - a, p["motif"], p["mode_paiement"],
            ])

        _creer_onglet(
            wb, "Paies Personnel", "PAIEMENTS AU PERSONNEL",
            ["N Paie", "Date", "Code", "Nom", "Prenom", "Fonction",
             "Salaire", "Avance", "Net", "Motif", "Mode"],
            lignes,
            largeurs=[16, 20, 14, 16, 16, 18, 14, 14, 14, 20, 16],
        )
        ws = wb["Paies Personnel"]
        for col in (7, 8, 9):
            _formater_nombre(ws, col)

        # ============ 5. AVANCES ============
        cursor.execute("""
            SELECT a.numero_avance, a.date_avance, p.code, p.nom, p.prenom,
                   a.montant, a.montant_deduit, a.motif
            FROM avances_personnel a
            JOIN personnel p ON a.personnel_id = p.id
            ORDER BY a.date_avance DESC
        """)
        avances = cursor.fetchall()

        lignes = []
        for a in avances:
            m = float(a["montant"] or 0)
            d = float(a["montant_deduit"] or 0)
            lignes.append([
                a["numero_avance"],
                a["date_avance"][:19] if a["date_avance"] else "",
                a["code"], a["nom"], a["prenom"],
                m, d, m - d, a["motif"] or "",
            ])

        _creer_onglet(
            wb, "Avances", "AVANCES SUR SALAIRE",
            ["N Avance", "Date", "Code", "Nom", "Prenom",
             "Montant", "Deduit", "Reste", "Motif"],
            lignes,
            largeurs=[16, 20, 14, 16, 16, 14, 14, 14, 25],
        )
        ws = wb["Avances"]
        for col in (6, 7, 8):
            _formater_nombre(ws, col)

        # ============ 6. DEPENSES ============
        cursor.execute("""
            SELECT date_depense, categorie, libelle, montant
            FROM depenses
            ORDER BY date_depense DESC
        """)
        depenses = cursor.fetchall()

        lignes = []
        for d in depenses:
            lignes.append([
                d["date_depense"][:19] if d["date_depense"] else "",
                d["categorie"], d["libelle"], float(d["montant"] or 0),
            ])

        _creer_onglet(
            wb, "Depenses", "DEPENSES GENERALES",
            ["Date", "Categorie", "Libelle", "Montant"],
            lignes,
            largeurs=[20, 20, 40, 16],
        )
        _formater_nombre(wb["Depenses"], 4)

        conn.close()

        # ============ 7. JOURNAL GENERAL ============
        try:
            annee = datetime.now().year
            journal = etat_journal(f"{annee}-01-01", f"{annee}-12-31")
            lignes = []
            for l in journal["lignes"]:
                lignes.append([
                    l["date"], l["journal"], l["piece"],
                    l["libelle"], l["compte"], l["intitule"],
                    float(l["debit"] or 0), float(l["credit"] or 0),
                ])
            _creer_onglet(
                wb, "Journal General", "JOURNAL GENERAL",
                ["Date", "Journal", "Piece", "Libelle", "Compte",
                 "Intitule", "Debit", "Credit"],
                lignes,
                largeurs=[12, 12, 16, 35, 10, 30, 14, 14],
            )
            ws = wb["Journal General"]
            _formater_nombre(ws, 7)
            _formater_nombre(ws, 8)
        except Exception as e:
            print(f"[Erreur journal] {e}")

        # ============ 8. BALANCE ============
        try:
            annee = datetime.now().year
            balance = etat_balance(f"{annee}-01-01", f"{annee}-12-31")
            lignes = []
            for l in balance["lignes"]:
                lignes.append([
                    l["compte"], l["intitule"],
                    float(l["mvt_debit"] or 0), float(l["mvt_credit"] or 0),
                    float(l["solde_debit"] or 0), float(l["solde_credit"] or 0),
                ])
            _creer_onglet(
                wb, "Balance", "BALANCE GENERALE",
                ["Compte", "Intitule", "Mvt Debit", "Mvt Credit",
                 "Solde Deb", "Solde Cred"],
                lignes,
                largeurs=[12, 35, 16, 16, 16, 16],
            )
            ws = wb["Balance"]
            for col in (3, 4, 5, 6):
                _formater_nombre(ws, col)
        except Exception as e:
            print(f"[Erreur balance] {e}")

        # ============ 9. BILAN ============
        try:
            annee = datetime.now().year
            bilan = etat_bilan(f"{annee}-01-01", f"{annee}-12-31")
            lignes = []
            max_rows = max(len(bilan["actif"]), len(bilan["passif"]))
            for i in range(max_rows):
                a = bilan["actif"][i] if i < len(bilan["actif"]) else {}
                p = bilan["passif"][i] if i < len(bilan["passif"]) else {}
                lignes.append([
                    a.get("compte", ""), a.get("intitule", ""),
                    float(a.get("montant", 0)),
                    p.get("compte", ""), p.get("intitule", ""),
                    float(p.get("montant", 0)),
                ])
            _creer_onglet(
                wb, "Bilan", "BILAN",
                ["Cpte Actif", "Actif", "Montant", "Cpte Passif", "Passif", "Montant"],
                lignes,
                largeurs=[12, 30, 16, 12, 30, 16],
            )
            ws = wb["Bilan"]
            _formater_nombre(ws, 3)
            _formater_nombre(ws, 6)
        except Exception as e:
            print(f"[Erreur bilan] {e}")

        # ============ 10. COMPTE DE RESULTAT ============
        try:
            annee = datetime.now().year
            cr = etat_compte_resultat(f"{annee}-01-01", f"{annee}-12-31")
            lignes = []
            for p in cr["produits"]:
                lignes.append(["PRODUIT", p["compte"], p["intitule"],
                               float(p["montant"])])
            for c in cr["charges"]:
                lignes.append(["CHARGE", c["compte"], c["intitule"],
                               float(c["montant"])])
            _creer_onglet(
                wb, "Compte Resultat", "COMPTE DE RESULTAT",
                ["Type", "Compte", "Intitule", "Montant"],
                lignes,
                largeurs=[12, 12, 40, 16],
            )
            _formater_nombre(wb["Compte Resultat"], 4)
        except Exception as e:
            print(f"[Erreur compte resultat] {e}")

        # ============ SAUVEGARDER ============
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        nom_fichier = f"b_ndeke_export_{timestamp}.xlsx"
        chemin = os.path.join(EXPORTS_DIR, nom_fichier)

        wb.save(chemin)
        return True, chemin

    except Exception as e:
        import traceback
        traceback.print_exc()
        return False, f"Erreur : {str(e)}"


def get_dossier_exports():
    return EXPORTS_DIR


def lister_exports():
    """Liste les fichiers Excel crees"""
    if not os.path.exists(EXPORTS_DIR):
        return []
    fichiers = []
    for f in os.listdir(EXPORTS_DIR):
        if f.endswith(".xlsx"):
            chemin = os.path.join(EXPORTS_DIR, f)
            fichiers.append({
                "nom": f,
                "chemin": chemin,
                "taille": os.path.getsize(chemin),
                "date": datetime.fromtimestamp(os.path.getmtime(chemin)),
            })
    return sorted(fichiers, key=lambda x: x["date"], reverse=True)